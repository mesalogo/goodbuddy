"""Real Electron/Docs right-click AI UI for Word, Excel and PowerPoint.
Conversation responses are controlled fixtures; full Runtime is tested separately.
"""
import argparse
import json
import os
from pathlib import Path
import shutil
import signal
import subprocess
import tempfile
import time
import zipfile
from playwright.sync_api import sync_playwright
from docx import Document
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from pptx import Presentation

parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument('--web', required=True)
parser.add_argument('--formats', nargs='+', choices=['docx','xlsx','pptx'], default=['docx','xlsx','pptx'])
args = parser.parse_args()
repo = Path(__file__).resolve().parents[1]
root = Path(tempfile.mkdtemp(prefix='office-ai-', dir=repo / 'temp'))
for source, output, platform in [
    ('src/main/office-service.ts', 'service.cjs', 'node'),
    ('src/main/office-file-worker.ts', 'office-file-worker.js', 'node'),
    ('src/preload/office-editor.ts', 'office-editor.cjs', 'node'),
    ('tests/support/office-desktop-panel.tsx', 'panel.js', 'browser')]:
    subprocess.run([str(repo / 'node_modules/.bin/esbuild'), source, '--bundle', '--platform='+platform,
                    '--format='+('cjs' if platform == 'node' else 'iife'), '--external:electron', '--jsx=automatic', '--outfile='+str(root/output)], cwd=repo, check=True)
shutil.copyfile(repo / 'tests/support/office-desktop.html', root / 'index.html')
(root / 'package.json').write_text('{"type":"commonjs"}')
doc = Document(); doc.add_paragraph('GB_WORD_ORIGINAL'); doc.save(root/'source.docx')
book = Workbook(); book.active['A1'] = 'GB_CELL_ORIGINAL'; book.active['A1'].font = Font(bold=True); book.active['B1'] = '=1+2'; book.active['B1'].number_format = '0.00'; book.save(root/'source.xlsx')
slides = Presentation(); slide = slides.slides.add_slide(slides.slide_layouts[0]); slide.shapes.title.text = 'GB_PPT_ORIGINAL'; slides.save(root/'source.pptx')
geometry = [(s.left,s.top,s.width,s.height) for s in slides.slides[0].shapes]
def contents(extension):
    with zipfile.ZipFile(root/f'source.{extension}') as archive:
        return '\n'.join(archive.read(n).decode('utf-8', errors='replace') for n in archive.namelist() if n.endswith('.xml'))

try:
    for extension in args.formats:
        reply = '[["GB_AI_CELL","=1+2"]]' if extension == 'xlsx' else 'GB_AI_REPLACEMENT'
        env = dict(os.environ, TMPDIR=str(root), GB_OFFICE_PROBE_ROOT=str(root), GB_OFFICE_PROBE_WEB=args.web,
                   GB_OFFICE_PROBE_FILENAME='source.'+extension, GB_OFFICE_PROBE_REPLY=reply)
        env.pop('ELECTRON_RUN_AS_NODE', None)
        with (root/f'{extension}.log').open('w') as log:
            child = subprocess.Popen(['xvfb-run','-a',str(repo/'node_modules/electron/dist/electron'),'--no-sandbox',str(repo/'tests/support/office-desktop-main.cjs')],env=env,stdout=log,stderr=log,start_new_session=True)
            try:
                with sync_playwright() as p:
                    for _ in range(60):
                        try: browser = p.chromium.connect_over_cdp('http://127.0.0.1:19222'); break
                        except Exception:
                            if child.poll() is not None: raise RuntimeError('Electron exited')
                            time.sleep(.3)
                    else: raise RuntimeError('CDP unavailable')
                    panel = browser.contexts[0].pages[0]
                    panel.get_by_role('button', name='选择 Office 文件并编辑', exact=True).click()
                    editor = None
                    for _ in range(120):
                        editor = next((page for page in browser.contexts[0].pages if '/office/editor?' in page.url), None)
                        if editor and editor.evaluate('document.documentElement.dataset.officeReady === "true"'): break
                        panel.wait_for_timeout(500)
                    else: raise RuntimeError('Editor not ready')
                    assert panel.get_by_text('AI 编辑选区', exact=True).count() == 0
                    frame = editor.frames[1]
                    if frame.get_by_role('button',name='取消',exact=True).count(): frame.get_by_role('button',name='取消',exact=True).click()
                    if extension == 'docx':
                        editor.mouse.click(650,560); editor.keyboard.press('Control+End'); editor.keyboard.press('Control+a')
                    elif extension == 'xlsx':
                        box = frame.locator('#ws-canvas').bounding_box()
                        editor.mouse.click(box['x']+90,box['y']+40)
                        editor.keyboard.press('Control+Home'); editor.keyboard.press('Shift+ArrowRight')
                    else:
                        editor.mouse.dblclick(800,475); editor.keyboard.press('Control+a')
                    for _ in range(40):
                        try:
                            selected = editor.evaluate('window.goodbuddyOffice.selection()')
                            break
                        except Exception as error:
                            if '尚未就绪' not in str(error): raise
                            editor.wait_for_timeout(250)
                    else: raise RuntimeError('GoodBuddy plugin did not initialize')
                    print('SELECTION', extension, selected, flush=True)
                    editor.screenshot(path=str(root/f'{extension}-selected.png'))
                    if extension == 'docx': editor.mouse.click(400,300,button='right')
                    elif extension == 'xlsx': editor.mouse.click(box['x']+50,box['y']+30,button='right')
                    else: editor.mouse.click(800,475,button='right')
                    menu = frame.get_by_text('AI 处理所选单元格…' if extension == 'xlsx' else 'AI 润色…', exact=True)
                    try: menu.wait_for(timeout=10000)
                    except Exception:
                        editor.screenshot(path=str(root/f'{extension}-menu-failed.png'))
                        raise
                    menu.click()
                    modal = editor.get_by_role('dialog')
                    modal.wait_for()
                    editor.screenshot(path=str(root/f'{extension}-dialog.png'))
                    print('DIALOG',modal.inner_text(),flush=True)
                    modal.get_by_label('修改要求',exact=True).fill('润色所选内容，保持格式与公式。')
                    modal.get_by_role('button',name='开始润色',exact=True).click()
                    result = modal.get_by_label('第 1 行第 1 列' if extension == 'xlsx' else '修改结果',exact=True)
                    result.wait_for(timeout=15000)
                    assert result.input_value() == ('GB_AI_CELL' if extension == 'xlsx' else reply)
                    if extension == 'xlsx':
                        formula = modal.get_by_label('第 1 行第 2 列',exact=True)
                        formula.fill('99')
                        modal.get_by_role('button',name='应用到原选区',exact=True).click()
                        modal.get_by_role('alert').filter(has_text='公式已受保护').wait_for()
                        formula.fill('=1+2')
                    modal.get_by_role('button',name='应用到原选区',exact=True).click()
                    modal.wait_for(state='hidden')
                    panel.evaluate('window.goodbuddy.probeSaveKey(false)')
                    editor.get_by_role('status').filter(has_text='已保存到文件').wait_for(timeout=65000)
                    assert ('GB_AI_CELL' if extension == 'xlsx' else reply) in contents(extension)
                    if extension == 'xlsx': assert '<f>1+2</f>' in contents(extension)
                    if extension == 'xlsx':
                        saved_sheet = load_workbook(root/'source.xlsx').active
                        assert saved_sheet['A1'].font.bold
                        assert saved_sheet['B1'].number_format == '0.00'
                    if extension == 'pptx':
                        saved_geometry = [(s.left,s.top,s.width,s.height) for s in Presentation(root/'source.pptx').slides[0].shapes]
                        assert len(saved_geometry) == len(geometry)
                        # OOXML round-trip can round integer EMUs (< 0.000003 inch).
                        assert all(abs(a-b) <= 2 for before,after in zip(geometry,saved_geometry) for a,b in zip(before,after))
                    print('PASS:', extension, 'right-click -> modal -> GB IPC conversation reply -> apply -> native save -> source verified', flush=True)
                    instance_id = panel.locator('[data-office-instance]').get_attribute('data-office-instance')
                    original_marker = {'docx':'GB_WORD_ORIGINAL','xlsx':'GB_CELL_ORIGINAL','pptx':'GB_PPT_ORIGINAL'}[extension]
                    for key, marker in [('Z',original_marker),('Y','GB_AI_CELL' if extension == 'xlsx' else reply)]:
                        icon = '.btn-undo' if key == 'Z' else '.btn-redo'
                        frame.locator('button:visible').filter(has=frame.locator(icon)).first.click()
                        for _ in range(20):
                            panel.evaluate('id => window.goodbuddy.office.save({instanceId:id})',instance_id)
                            if marker in contents(extension): break
                            panel.wait_for_timeout(500)
                        else: raise RuntimeError(f'{extension}: undo/redo failed: {key}')
                    print('PASS:',extension,'single native undo/redo preserved content',flush=True)
                    assert panel.evaluate('id => window.goodbuddy.office.close({instanceId:id})',instance_id)
                    panel.evaluate('window.goodbuddy.quitProbe()'); browser.close(); child.wait(timeout=10)
            finally:
                if child.poll() is None: os.killpg(child.pid,signal.SIGTERM); child.wait(timeout=10)
    print('PASS: selected Office AI dialogs:', ','.join(args.formats), flush=True)
except Exception:
    print('Diagnostics:',root,flush=True)
    raise
else:
    shutil.rmtree(root)
